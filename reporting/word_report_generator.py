"""
Word Report Generator
=====================
Converts Report Agent output into a formatted Microsoft Word report.
This module renders report content; it does not calculate sales analytics.
"""
import json
from groq import Groq

import os
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# ---------------------------------------------------------------------
# General helper functions
# ---------------------------------------------------------------------

def _dict(value: Any) -> Dict[str, Any]:
    """Safely convert an object to a dictionary."""
    if isinstance(value, dict):
        return value

    if value is None:
        return {}

    # DataFrames are handled separately as tabular data.
    if hasattr(value, "columns") and hasattr(value, "to_dict"):
        return {}

    if hasattr(value, "to_dict"):
        try:
            result = value.to_dict()
            if isinstance(result, dict):
                return result
        except Exception:
            pass

    if hasattr(value, "__dict__"):
        try:
            return vars(value)
        except Exception:
            pass

    return {}


def _list(value: Any) -> List[Any]:
    """Convert a value into a list without discarding existing lists."""
    if value is None:
        return []

    if isinstance(value, list):
        return value

    if isinstance(value, tuple):
        return list(value)

    return [value]


def _text(value: Any, default: str = "N/A") -> str:
    """Convert values into readable text without raw nested structures."""
    if value is None:
        return default

    try:
        if value != value:
            return "—"
    except Exception:
        pass

    if hasattr(value, "strftime") and callable(value.strftime):
        try:
            return value.strftime("%d %b %Y")
        except Exception:
            pass

    if isinstance(value, dict):
        parts = []

        for key, item in value.items():
            if isinstance(item, (dict, list, tuple)):
                continue

            parts.append(
                f"{str(key).replace('_', ' ').title()}: "
                f"{_text(item, '')}"
            )

        return "; ".join(parts) if parts else "Details available"

    if isinstance(value, (list, tuple)):
        return "; ".join(_text(item, "") for item in value)

    if isinstance(value, float):
        if value != value:
            return "—"

        if value.is_integer():
            return f"{int(value):,}"

        return f"{value:,.2f}"

    return str(value)


def _display_cell(key: Any, value: Any) -> str:
    """Format table values according to their column name."""
    name = str(key).lower().replace("_", " ")

    if value is None:
        return "—"

    try:
        if value != value:
            return "—"
    except Exception:
        pass

    if hasattr(value, "strftime") and callable(value.strftime):
        try:
            return value.strftime("%d %b %Y")
        except Exception:
            pass

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)

        if any(
            token in name
            for token in (
                "percent", "percentage", "growth",
                "change pct", "mape",
            )
        ):
            return f"{number:,.2f}%"

        if any(
            token in name
            for token in (
                "horizon", "count", "rows", "columns",
                "quantity", "orders", "customers",
                "checks", "observations", "rank",
            )
        ):
            if number.is_integer():
                return f"{int(number):,}"

            return f"{number:,.2f}"

        if any(
            token in name
            for token in (
                "sales", "revenue", "amount", "price",
                "value", "mae", "rmse", "forecast",
            )
        ):
            if number.is_integer():
                return f"₹{number:,.0f}"

            return f"₹{number:,.2f}"

        if number.is_integer():
            return f"{int(number):,}"

        return f"{number:,.2f}"

    return _text(value, "")


def _has_data(value: Any) -> bool:
    """Check whether a value contains data, including pandas objects."""
    if value is None:
        return False

    if hasattr(value, "empty"):
        try:
            return not bool(value.empty)
        except Exception:
            return False

    if isinstance(value, (dict, list, tuple, set, str, bytes)):
        return len(value) > 0

    try:
        return bool(value)
    except (TypeError, ValueError):
        return True


def _first(
    data: Dict[str, Any],
    *keys: str,
    default: Any = None,
) -> Any:
    """Return the first existing non-None value."""
    for key in keys:
        if key in data and data[key] is not None:
            return data[key]

    return default


def _fmt(value: Any, decimals: int = 2) -> str:
    """Format a number with thousands separators."""
    if value is None:
        return "N/A"

    try:
        number = float(value)

        if number.is_integer():
            return f"{int(number):,}"

        return f"{number:,.{decimals}f}"

    except (TypeError, ValueError):
        return _text(value)


def _normalise_report(report: Any) -> Dict[str, Any]:
    return _dict(report)


def _normalise_section(section: Any) -> Dict[str, Any]:
    return _dict(section)


def _find_section(
    report: Dict[str, Any],
    section_id: str,
    title_contains: Optional[str] = None,
) -> Dict[str, Any]:
    """Find a report section by ID or partial title."""
    sections = _list(report.get("sections"))

    for section in sections:
        item = _normalise_section(section)

        if item.get("section_id") == section_id:
            return item

    if title_contains:
        phrase = title_contains.lower()

        for section in sections:
            item = _normalise_section(section)
            title = _text(item.get("title"), "").lower()

            if phrase in title:
                return item

    return {}


def _content(section: Dict[str, Any]) -> Any:
    return section.get("content", {}) if section else {}


def _as_records(value: Any) -> List[Dict[str, Any]]:
    """Convert lists, dictionaries and DataFrames into table records."""
    if value is None:
        return []

    if hasattr(value, "to_dict") and hasattr(value, "columns"):
        try:
            records = value.to_dict(orient="records")
            return [
                row for row in records
                if isinstance(row, dict)
            ]
        except Exception:
            pass

    if hasattr(value, "to_dict") and not isinstance(value, dict):
        try:
            converted = value.to_dict()

            if isinstance(converted, dict):
                return [
                    {"Metric": key, "Value": item}
                    for key, item in converted.items()
                ]
        except Exception:
            pass

    if isinstance(value, dict):
        return [
            {"Metric": key, "Value": item}
            for key, item in value.items()
        ]

    records = []

    for item in _list(value):
        if isinstance(item, dict):
            records.append(item)

        elif isinstance(item, (list, tuple)):
            records.append({
                f"Column {index + 1}": item_value
                for index, item_value in enumerate(item)
            })

        else:
            records.append({"Item": item})

    return records


def _set_cell_shading(cell, fill: str = "1F4E78") -> None:
    """Apply background shading to a Word table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))

    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)

    shading.set(qn("w:fill"), fill)


def _format_table(table, header: bool = True) -> None:
    """Apply consistent formatting to a Word table."""
    table.style = "Table Grid"

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)

                for run in paragraph.runs:
                    run.font.size = Pt(9)

        if header and row_index == 0:
            for cell in row.cells:
                _set_cell_shading(cell)

                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def _add_heading(
    document: Document,
    title: str,
    level: int = 1,
) -> None:
    heading = document.add_heading(title, level=level)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(5)


def _add_note(document: Document, message: str) -> None:
    paragraph = document.add_paragraph(message)
    paragraph.paragraph_format.space_after = Pt(5)
    # ---------------------------------------------------------------------
# Table and title-page helpers
# ---------------------------------------------------------------------

def _add_key_value_table(
    document: Document,
    rows: List[tuple],
) -> None:
    if not rows:
        _add_note(document, "No details were provided for this section.")
        return

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = _text(key)
        cells[1].text = _display_cell(key, value)

    _format_table(table)


def _add_records_table(
    document: Document,
    records: Any,
    preferred_keys=None,
) -> None:
    rows = _as_records(records)

    if not rows:
        _add_note(document, "No records were provided for this section.")
        return

    if preferred_keys is None:
        preferred_keys = []

    keys = []

    for key in preferred_keys:
        if any(key in row for row in rows):
            keys.append(key)

    for row in rows:
        for key in row:
            if key not in keys:
                keys.append(key)

    keys = keys[:8]

    table = document.add_table(
        rows=1,
        cols=max(1, len(keys)),
    )

    for index, key in enumerate(keys):
        table.rows[0].cells[index].text = (
            str(key).replace("_", " ").title()
        )

    for row in rows[:100]:
        cells = table.add_row().cells

        for index, key in enumerate(keys):
            cells[index].text = _display_cell(
                key,
                row.get(key, ""),
            )

    _format_table(table)


def _add_title_page(
    document: Document,
    report: Dict[str, Any],
    dataset_name: Optional[str] = None,
) -> None:
    document.add_paragraph()

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        _text(
            report.get("title"),
            "AI-Powered Sales Intelligence Report",
        )
    )
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(31, 78, 121)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "AI-Powered Multi-Agent Sales Intelligence System"
    )
    run.font.size = Pt(14)

    if dataset_name:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"Dataset: {dataset_name}").bold = True

    report_type = report.get("report_type")

    if report_type:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(
            f"Report Type: {_text(report_type)}"
        )

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Generated by the Report Agent").italic = True

    document.add_page_break()


def _add_executive_summary(
    document: Document,
    report: Dict[str, Any],
) -> None:
    _add_heading(document, "1. Executive Summary")

    summary = report.get("executive_summary")

    if not _has_data(summary):
        summary = report.get("summary")

    if not _has_data(summary):
        summary = (
            "This report consolidates validated outputs from the "
            "AI-powered multi-agent sales intelligence workflow. "
            "Interpret the results in the context of the uploaded "
            "dataset and the limitations listed in this report."
        )

    _add_note(document, _text(summary))


def _add_dataset_overview(
    document: Document,
    content: Any,
) -> None:
    _add_heading(document, "2. Dataset Overview")
    data = _dict(content)

    if not _has_data(data):
        _add_note(
            document,
            "Dataset overview information was not available.",
        )
        return

    rows = []

    aliases = [
        ("Dataset Name", ("dataset_name", "name")),
        ("Records", ("rows", "record_count", "records", "row_count")),
        ("Columns", ("columns", "column_count", "number_of_columns")),
        ("Missing Values", ("missing_values", "missing_count")),
        ("Duplicate Rows", ("duplicate_rows", "duplicates", "duplicate_count")),
    ]

    used = {
        key
        for _, keys in aliases
        for key in keys
    }

    for label, keys in aliases:
        value = _first(data, *keys)

        if value is not None:
            rows.append((label, value))

    for key, value in data.items():
        if (
            key not in used
            and not isinstance(value, (dict, list, tuple))
        ):
            rows.append((
                str(key).replace("_", " ").title(),
                value,
            ))

    _add_key_value_table(document, rows)


def _add_data_cleaning(
    document: Document,
    content: Any,
) -> None:
    _add_heading(document, "3. Data Quality and Cleaning")

    data = _dict(content)
    cleaning = _dict(data.get("cleaning_result"))

    validation = _dict(
        _first(data, "cleaning_validation", "validation", default={})
    )

    rows = []

    fields = [
        (
            "Cleaning Status",
            cleaning,
            ("status", "cleaning_status", "overall_status"),
        ),
        (
            "Issues Detected",
            cleaning,
            ("issues_detected", "total_issues", "issue_count"),
        ),
        (
            "Issues Resolved",
            cleaning,
            ("issues_resolved", "resolved_count"),
        ),
        (
            "Validation Status",
            validation,
            ("overall_status", "status", "validation_status"),
        ),
        (
            "Checks Passed",
            validation,
            ("passed", "passed_checks", "checks_passed"),
        ),
        (
            "Checks Failed",
            validation,
            ("failed", "failed_checks", "checks_failed"),
        ),
    ]

    for label, source, keys in fields:
        value = _first(source, *keys)

        if value is not None:
            rows.append((label, value))

    if rows:
        _add_key_value_table(document, rows)
    else:
        _add_note(
            document,
            "Cleaning or validation details were not supplied.",
        )

    issues = _first(
        cleaning,
        "issues",
        "actions",
        "cleaning_actions",
        default=[],
    )

    if _has_data(issues):
        document.add_paragraph(
            "Cleaning Details",
            style="Heading 2",
        )
        _add_records_table(document, issues)


def _add_trends(
    document: Document,
    content: Any,
) -> None:
    _add_heading(document, "4. Sales Trends")

    if isinstance(content, list):
        data = {"items": content}
    else:
        data = _dict(content)

    observations = _first(
        data,
        "items",
        "data",
        "trend_data",
        "series",
        "results",
        "monthly",
        "monthly_sales",
        "monthly_trends",
        "trends",
        default=[],
    )

    summary_rows = []

    summary_fields = [
        ("Trend Summary", ("summary", "description", "trend_summary", "interpretation")),
        ("Direction", ("direction", "trend_direction")),
        ("Start Period", ("start_period", "start_date", "start")),
        ("End Period", ("end_period", "end_date", "end")),
        ("Observations", ("observations", "observation_count", "count")),
    ]

    for label, keys in summary_fields:
        value = _first(data, *keys)

        if (
            value is not None
            and not isinstance(value, (dict, list, tuple))
        ):
            summary_rows.append((label, value))

    if summary_rows:
        _add_key_value_table(document, summary_rows)

    if _has_data(observations):
        _add_records_table(
            document,
            observations,
            preferred_keys=[
                "period",
                "month",
                "date",
                "label",
                "sales",
                "value",
                "total_sales",
                "change_pct",
                "percentage_change",
                "growth",
            ],
        )
    elif not summary_rows:
        _add_note(document, "No trend observations were supplied.")


def _add_rankings(
    document: Document,
    content: Any,
) -> None:
    _add_heading(
        document,
        "5. Sales and Performance Rankings",
    )

    data = _dict(content)

    groups = [
        (
            "Product Sales",
            ("product_sales", "product_categories", "top_products"),
        ),
        (
            "Product Quantity",
            ("product_quantity", "quantity_by_product", "top_quantity_products"),
        ),
        (
            "Regional Sales",
            ("regional_sales", "regions", "sales_by_region", "city_sales"),
        ),
        (
            "Customer Sales",
            ("customer_sales", "customers", "top_customers"),
        ),
    ]

    found = False

    for heading, keys in groups:
        value = _first(data, *keys)

        if _has_data(value):
            document.add_paragraph(
                heading,
                style="Heading 2",
            )

            _add_records_table(
                document,
                value,
                preferred_keys=[
                    "rank",
                    "product",
                    "product_name",
                    "category",
                    "region",
                    "city",
                    "customer",
                    "customer_id",
                    "sales",
                    "total_sales",
                    "quantity",
                    "value",
                ],
            )

            found = True

    if not found:
        items = _first(
            data,
            "items",
            "rankings",
            "data",
            "results",
            default=None,
        )

        if items is None and isinstance(content, list):
            items = content

        if _has_data(items):
            _add_records_table(
                document,
                items,
                preferred_keys=[
                    "rank",
                    "name",
                    "product",
                    "category",
                    "region",
                    "customer",
                    "sales",
                    "total_sales",
                    "quantity",
                    "value",
                ],
            )
            found = True

    if not found:
        _add_note(document, "No ranking data was supplied.")
# ---------------------------------------------------------------------
# KPI rendering and semantic number formatting
# ---------------------------------------------------------------------

def _add_kpis(
    document: Document,
    content: Any,
) -> None:
    """Render KPIs with formatting based on each metric's meaning."""
    _add_heading(
        document,
        "6. Key Performance Indicators (KPIs)",
    )

    data = _dict(content)

    if isinstance(content, list):
        metrics = content
    else:
        metrics = _first(
            data,
            "kpis",
            "kpi_results",
            "metrics",
            "items",
            "results",
            default=None,
        )

    # Handle KPIs supplied as a list of dictionaries.
    if _has_data(metrics):
        records = _as_records(metrics)

        if records:
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Metric"
            table.rows[0].cells[1].text = "Value"

            for record in records:
                metric_name = _first(
                    record,
                    "name",
                    "metric",
                    "kpi",
                    "label",
                    default="Metric",
                )

                value = _first(
                    record,
                    "value",
                    "result",
                    "amount",
                    default=None,
                )

                # Support records using a different value-field name.
                if value is None:
                    remaining = [
                        (key, val)
                        for key, val in record.items()
                        if key not in {
                            "name",
                            "metric",
                            "kpi",
                            "label",
                            "description",
                            "unit",
                        }
                    ]

                    if remaining:
                        value = remaining[0][1]

                formatted_value = _format_kpi_value(
                    metric_name,
                    value,
                    record.get("unit"),
                )

                cells = table.add_row().cells
                cells[0].text = _text(metric_name)
                cells[1].text = formatted_value

            _format_table(table)
            return

    # Handle KPI results supplied directly as a dictionary.
    excluded = {
        "status",
        "message",
        "description",
        "validation",
        "metrics",
        "metadata",
    }

    direct = [
        (key, value)
        for key, value in data.items()
        if key not in excluded
        and not isinstance(value, (dict, list, tuple))
    ]

    if direct:
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Metric"
        table.rows[0].cells[1].text = "Value"

        for metric_name, value in direct:
            cells = table.add_row().cells

            cells[0].text = str(metric_name).replace(
                "_", " "
            ).title()

            cells[1].text = _format_kpi_value(
                metric_name,
                value,
            )

        _format_table(table)

    else:
        _add_note(document, "No KPI values were supplied.")



def _format_kpi_value(
    metric_name: Any,
    value: Any,
    unit: Any = None,
) -> str:
    """Format KPI values according to their business meaning."""

    name = str(metric_name).strip().lower().replace("_", " ")
    unit_text = str(unit).strip().lower() if unit is not None else ""

    if value is None:
        return "—"

    try:
        number = float(value)

        if number != number:  # NaN
            return "—"

    except (TypeError, ValueError):
        return _text(value)

    # Check monetary metrics BEFORE counts.
    currency_terms = (
        "sales",
        "revenue",
        "amount",
        "turnover",
        "order value",
        "average order value",
        "selling price",
        "unit price",
        "profit",
        "cost",
        "income",
    )

    is_currency = (
        any(term in name for term in currency_terms)
        or unit_text in {"₹", "inr", "rs", "rupee", "rupees"}
    )

    if is_currency:
        if number.is_integer():
            return f"₹{number:,.0f}"
        return f"₹{number:,.2f}"

    # Counts and quantities must not receive a currency symbol.
    count_terms = (
        "order",
        "customer",
        "quantity",
        "units",
        "record",
        "row count",
        "column count",
        "transaction count",
        "number of",
    )

    if any(term in name for term in count_terms):
        if number.is_integer():
            return f"{number:,.0f}"
        return f"{number:,.2f}"

    # Format percentages.
    if (
        "percent" in name
        or "percentage" in name
        or unit_text in {"%", "percent"}
    ):
        return f"{number:,.2f}%"

    # Default numeric formatting.
    if number.is_integer():
        return f"{number:,.0f}"

    return f"{number:,.2f}"


# ---------------------------------------------------------------------
# Visualisation rendering
# ---------------------------------------------------------------------

def _add_visualizations(
    document: Document,
    content: Any,
) -> None:
    _add_heading(document, "7. Visualizations")

    data = _dict(content)

    charts = _first(
        data,
        "charts",
        "visualizations",
        "items",
        "figures",
        default=[],
    )

    if isinstance(content, list):
        charts = content

    charts = _list(charts)
    added = 0

    for chart in charts:
        item = _dict(chart)

        title = _first(
            item,
            "title",
            "name",
            "chart_title",
            default="Chart",
        )

        description = _first(
            item,
            "description",
            "caption",
            "summary",
            default=None,
        )

        path = _first(
            item,
            "path",
            "image_path",
            "file_path",
            "chart_path",
            "saved_path",
        )

        document.add_paragraph(
            _text(title),
            style="Heading 2",
        )

        if _has_data(description):
            _add_note(document, _text(description))

        if _has_data(path):
            candidates = [os.path.abspath(str(path))]

            if not os.path.isabs(str(path)):
                candidates.append(
                    os.path.abspath(
                        os.path.join(os.getcwd(), str(path))
                    )
                )

            image_path = next(
                (
                    candidate
                    for candidate in candidates
                    if os.path.isfile(candidate)
                ),
                None,
            )

            if image_path:
                try:
                    document.add_picture(
                        image_path,
                        width=Inches(6.2),
                    )
                    added += 1

                except Exception as exc:
                    _add_note(
                        document,
                        f"Could not embed chart image: {exc}",
                    )

            else:
                _add_note(
                    document,
                    f"Chart image file was not found: {path}",
                )

        elif not _has_data(description):
            _add_note(
                document,
                "Chart metadata was provided without an image path.",
            )

    if not _has_data(charts):
        _add_note(
            document,
            "No visualization metadata was supplied.",
        )

    elif added:
        _add_note(
            document,
            f"Embedded {added} chart image(s).",
        )


# ---------------------------------------------------------------------
# Forecast rendering
# ---------------------------------------------------------------------

def _add_forecast(
    document: Document,
    content: Any,
) -> None:
    """Render forecast evaluation and values as readable tables."""
    _add_heading(document, "8. Sales Forecast")

    data = _dict(content)

    if not _has_data(data):
        _add_note(
            document,
            "Forecast results were not available for this dataset.",
        )
        return

    status = _first(
        data,
        "status",
        "forecast_status",
    )

    eligible = data.get("eligible")

    if (
        eligible is False
        or str(status).lower() in {
            "unavailable",
            "ineligible",
            "skipped",
            "failed",
        }
    ):
        _add_note(
            document,
            _text(
                _first(
                    data,
                    "reason",
                    "message",
                    default="Forecasting was unavailable.",
                )
            ),
        )
        return

    evaluation = _dict(
        _first(
            data,
            "evaluation",
            "model_evaluation",
            "evaluation_metrics",
            "metrics",
            default={},
        )
    )

    nested_results = _dict(
        _first(
            evaluation,
            "results",
            "model_results",
            "models",
            default={},
        )
    )

    if (
        not nested_results
        and evaluation
        and all(
            isinstance(value, dict)
            for value in evaluation.values()
        )
    ):
        nested_results = evaluation

    selected_model = _first(
        data,
        "best_model",
        "selected_model",
        "model_name",
        "model",
    )

    if not selected_model:
        selected_model = _first(
            evaluation,
            "best_model",
            "selected_model",
            "model_name",
        )

    best_mae = _first(
        data,
        "best_mae",
        "mae",
        "mean_absolute_error",
    )

    if best_mae is None:
        best_mae = _first(
            evaluation,
            "best_mae",
            "mae",
        )

    selected_eval = {}

    if selected_model and nested_results:
        target = str(selected_model).strip().lower()

        for name, metrics in nested_results.items():
            metric_data = _dict(metrics)

            model_name = str(
                _first(
                    metric_data,
                    "model_name",
                    default="",
                )
            ).strip().lower()

            if (
                str(name).strip().lower() == target
                or model_name == target
            ):
                selected_eval = metric_data
                break

        if not selected_eval:
            for name, metrics in nested_results.items():
                if (
                    target in str(name).lower()
                    or str(name).lower() in target
                ):
                    selected_eval = _dict(metrics)
                    break

    summary_rows = []

    if status is not None:
        summary_rows.append(("Forecast Status", status))

    if selected_model:
        summary_rows.append(("Selected Model", selected_model))

    frequency = _first(
        data,
        "frequency",
        default=_first(evaluation, "frequency"),
    )

    if frequency is not None:
        summary_rows.append(
            ("Frequency", str(frequency).upper())
        )

    horizon = _first(
        data,
        "horizon",
        "forecast_horizon",
    )

    forecast_values = _first(
        data,
        "forecast_values",
        "forecast",
        "predictions",
        "forecast_data",
        default=[],
    )

    if isinstance(forecast_values, dict):
        model_outputs = []

        for model_name, output in forecast_values.items():
            if isinstance(output, dict):
                item = dict(output)
                item.setdefault("model_name", model_name)
            else:
                item = {
                    "model_name": model_name,
                    "forecast": output,
                }

            model_outputs.append(item)

    else:
        model_outputs = _list(forecast_values)

    selected_output = None

    if selected_model:
        target = str(selected_model).strip().lower()

        for output in model_outputs:
            item = _dict(output)

            name = str(
                _first(
                    item,
                    "model_name",
                    "model",
                    "name",
                    default="",
                )
            ).strip().lower()

            if (
                name == target
                or (
                    name
                    and (name in target or target in name)
                )
            ):
                selected_output = item
                break

    if selected_output is None and len(model_outputs) == 1:
        selected_output = _dict(model_outputs[0])

    if (
        selected_output is None
        and model_outputs
        and all(
            isinstance(item, dict)
            and any(
                key in item
                for key in ("date", "period", "ds")
            )
            for item in model_outputs
        )
    ):
        selected_output = {
            "forecast": model_outputs,
        }

    selected_forecast = _first(
        _dict(selected_output),
        "forecast",
        "forecast_values",
        "predictions",
        "values",
        default=[],
    )

    if horizon is None and _has_data(selected_forecast):
        horizon = len(_list(selected_forecast))

    if horizon is not None:
        summary_rows.append(("Forecast Horizon", horizon))

    mae = _first(
        selected_eval,
        "mae",
        "mean_absolute_error",
        default=best_mae,
    )

    rmse = _first(
        selected_eval,
        "rmse",
        "root_mean_squared_error",
    )

    mape = _first(
        selected_eval,
        "mape",
        "mean_absolute_percentage_error",
    )

    for label, value in (
        ("MAE", mae),
        ("RMSE", rmse),
        ("MAPE", mape),
    ):
        if value is not None:
            summary_rows.append((label, value))

    message = _first(data, "reason", "message")

    if (
        message
        and not isinstance(message, (dict, list, tuple))
    ):
        summary_rows.append(("Notes", message))

    if summary_rows:
        _add_key_value_table(document, summary_rows)

    forecast_rows = []

    for record in _list(selected_forecast):
        row = _dict(record)

        if row:
            date_value = _first(
                row,
                "date",
                "period",
                "ds",
                "timestamp",
            )

            predicted = _first(
                row,
                "forecast",
                "prediction",
                "yhat",
                "value",
                "predicted_sales",
            )

            if date_value is not None or predicted is not None:
                forecast_rows.append({
                    "Forecast Date": date_value,
                    "Predicted Sales": predicted,
                })

        elif isinstance(record, (list, tuple)) and len(record) >= 2:
            forecast_rows.append({
                "Forecast Date": record[0],
                "Predicted Sales": record[1],
            })

    if forecast_rows:
        document.add_paragraph(
            "Forecast Values",
            style="Heading 2",
        )

        _add_records_table(
            document,
            forecast_rows,
            preferred_keys=[
                "Forecast Date",
                "Predicted Sales",
            ],
        )

    elif _has_data(forecast_values):
        _add_note(
            document,
            "Forecast values were supplied, but no forecast rows "
            "could be matched to the selected model.",
        )

    elif not summary_rows:
        _add_note(
            document,
            "Forecast section exists, but no displayable forecast "
            "values were supplied.",
        )
# ---------------------------------------------------------------------
# Business insights
# ---------------------------------------------------------------------

def _add_insights(
    document: Document,
    content: Any,
) -> None:
    _add_heading(document, "9. Business Insights")

    if isinstance(content, list):
        insights = content
    else:
        data = _dict(content)

        insights = _first(
            data,
            "insights",
            "key_insights",
            "business_insights",
            "items",
            "results",
            default=[],
        )

    insights = _list(insights)

    if not _has_data(insights):
        _add_note(document, "No business insights were supplied.")
        return

    for index, insight in enumerate(insights, start=1):
        item = _dict(insight)

        if not item:
            _add_note(
                document,
                f"{index}. {_text(insight)}",
            )
            continue

        title = _first(
            item,
            "title",
            "insight_title",
            "name",
            default=f"Insight {index}",
        )

        _add_heading(
            document,
            f"{index}. {_text(title)}",
            level=2,
        )

        statement = _first(
            item,
            "insight",
            "statement",
            "description",
            "text",
            "finding",
            default="",
        )

        if _has_data(statement):
            _add_note(document, _text(statement))

        evidence = _first(
            item,
            "evidence",
            "supporting_evidence",
            "data",
            "metric",
        )

        if evidence is not None:
            paragraph = document.add_paragraph()
            paragraph.add_run("Evidence: ").bold = True
            paragraph.add_run(_text(evidence))

        confidence = item.get("confidence")

        if confidence is not None:
            paragraph = document.add_paragraph()
            paragraph.add_run("Confidence: ").bold = True
            paragraph.add_run(_text(confidence))


# ---------------------------------------------------------------------
# Business recommendations
# ---------------------------------------------------------------------

def _add_recommendations(
    document: Document,
    content: Any,
) -> None:
    _add_heading(document, "10. Business Recommendations")

    if isinstance(content, list):
        recommendations = content
    else:
        data = _dict(content)

        recommendations = _first(
            data,
            "recommendations",
            "business_recommendations",
            "items",
            "results",
            default=[],
        )

    recommendations = _list(recommendations)

    if not _has_data(recommendations):
        _add_note(
            document,
            "No business recommendations were supplied.",
        )
        return

    for index, recommendation in enumerate(
        recommendations,
        start=1,
    ):
        item = _dict(recommendation)

        if not item:
            _add_note(
                document,
                f"{index}. {_text(recommendation)}",
            )
            continue

        title = _first(
            item,
            "title",
            "name",
            "recommendation_title",
            default=f"Recommendation {index}",
        )

        _add_heading(
            document,
            f"{index}. {_text(title)}",
            level=2,
        )

        action = _first(
            item,
            "recommendation",
            "action",
            "description",
            "text",
            default="",
        )

        if _has_data(action):
            _add_note(document, _text(action))

        optional_fields = [
            (
                "Priority",
                ("priority", "urgency"),
            ),
            (
                "Business Impact",
                (
                    "business_implication",
                    "implication",
                    "expected_impact",
                ),
            ),
            (
                "Evidence",
                (
                    "evidence",
                    "supporting_evidence",
                    "rationale",
                ),
            ),
        ]

        for label, keys in optional_fields:
            value = _first(item, *keys)

            if value is not None:
                paragraph = document.add_paragraph()
                paragraph.add_run(f"{label}: ").bold = True
                paragraph.add_run(_text(value))


# ---------------------------------------------------------------------
# Conclusion and limitations
# ---------------------------------------------------------------------


def _generate_conclusion_with_groq(
    report: Dict[str, Any],
) -> Optional[str]:
    """
    Generate an evidence-based conclusion from the report's
    actual analytics using the Groq API.
    """
    api_key = os.getenv("GROQ_API_KEY")

    if not api_key:
        return None

    # Extract only the sections relevant to the conclusion.
    relevant_sections = {}

    section_ids = [
        "dataset_overview",
        "data_quality",
        "kpis",
        "trends",
        "rankings",
        "forecast",
        "insights",
        "recommendations",
        "limitations",
    ]

    for section_id in section_ids:
        section = _find_section(report, section_id)

        if section:
            relevant_sections[section_id] = _content(section)

    evidence = {
        "report_title": report.get("title"),
        "report_type": report.get("report_type"),
        "executive_summary": report.get("executive_summary"),
        "sections": relevant_sections,
        "limitations": report.get("limitations"),
    }

    # Avoid sending an excessively large payload to the API.
    try:
        evidence_json = json.dumps(
            evidence,
            default=str,
            ensure_ascii=False,
        )

        evidence_json = evidence_json[:15000]

        model = os.getenv(
            "GROQ_MODEL",
            "llama-3.3-70b-versatile",
        )

        client = Groq(api_key=api_key)

        response = client.chat.completions.create(
            model=model,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a senior business intelligence "
                        "analyst writing the conclusion of a formal "
                        "sales analytics report. Use only the supplied "
                        "evidence. Do not invent numbers, trends, "
                        "causal explanations, business outcomes, or "
                        "model performance. Explain the most important "
                        "business findings, their implications, the "
                        "recommended direction, and relevant data or "
                        "forecast limitations. If evidence is missing, "
                        "state the limitation rather than guessing. "
                        "Write 2 to 4 coherent professional paragraphs "
                        "in clear business English. Do not use a title, "
                        "bullet points, or repeat the entire report."
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        "Write a specific, evidence-based conclusion "
                        "for this sales intelligence report.\n\n"
                        "Report evidence:\n"
                        f"{evidence_json}"
                    ),
                },
            ],
            temperature=0.2,
            max_tokens=500,
        )

        conclusion = (
            response.choices[0].message.content or ""
        ).strip()

        if conclusion and len(conclusion) >= 100:
            return conclusion

        return None

    except Exception as exc:
        print(
            "Groq conclusion generation failed: "
            f"{type(exc).__name__}: {exc}"
        )
        return None


def _add_conclusion(
    document: Document,
    report: Dict[str, Any],
) -> None:
    """Add an AI-generated conclusion with a reliable fallback."""
    _add_heading(document, "11. Conclusion")

    # Respect a conclusion already generated by an upstream agent.
    summary = _first(
        report,
        "conclusion",
        "closing_summary",
    )

    if _has_data(summary):
        _add_note(document, _text(summary))
        return

    # Generate a new conclusion from the available report evidence.
    conclusion = _generate_conclusion_with_groq(report)

    if conclusion:
        _add_note(document, conclusion)
        return

    # Fallback: retain report generation if the API is unavailable.
    insight_section = _find_section(
        report,
        "insights",
        "insight",
    )

    recommendation_section = _find_section(
        report,
        "recommendations",
        "recommendation",
    )

    def count_items(content: Any, keys: tuple) -> int:
        if isinstance(content, list):
            return len(content)

        data = _dict(content)
        items = _first(data, *keys, default=[])

        return len(_list(items))

    insight_count = count_items(
        _content(insight_section),
        ("insights", "key_insights", "items", "results"),
    )

    recommendation_count = count_items(
        _content(recommendation_section),
        ("recommendations", "business_recommendations",
         "items", "results"),
    )

    _add_note(
        document,
        (
            "The report consolidates "
            f"{insight_count} business insight(s) and "
            f"{recommendation_count} recommendation(s). "
            "Review these findings alongside the reported KPIs, "
            "sales trends, data-quality results and forecast "
            "limitations before making business decisions."
        ),
    )


def _add_limitations(
    document: Document,
    report: Dict[str, Any],
) -> None:
    _add_heading(document, "12. Limitations")

    section = _find_section(
        report,
        "limitations",
        "limitation",
    )

    content = _content(section)

    if isinstance(content, list):
        items = content
    else:
        data = _dict(content)

        items = _first(
            data,
            "limitations",
            "items",
            "warnings",
            "notes",
        )

    if _has_data(items):
        for item in _list(items):
            if isinstance(item, dict):
                item = _first(
                    item,
                    "description",
                    "message",
                    "limitation",
                    "text",
                    default=item,
                )

            document.add_paragraph(
                _text(item),
                style="List Bullet",
            )

        return

    report_limitations = report.get("limitations")

    if _has_data(report_limitations):
        for item in _list(report_limitations):
            document.add_paragraph(
                _text(item),
                style="List Bullet",
            )
    else:
        _add_note(
            document,
            "Forecasts and recommendations depend on data quality, "
            "available historical observations, column mapping and "
            "the assumptions used by the analytical workflow. "
            "Business users should review the findings before "
            "making consequential decisions.",
        )


def _add_technical_appendix(
    document: Document,
    report: Dict[str, Any],
) -> None:
    _add_heading(document, "Appendix: Report Metadata")

    rows = []

    fields = [
        ("Report Status", ("status", "report_status")),
        ("Report Type", ("report_type",)),
        ("Generated At", ("generated_at", "created_at", "timestamp")),
        ("Dataset", ("dataset_name",)),
        ("Workflow Status", ("workflow_status",)),
    ]

    for label, keys in fields:
        value = _first(report, *keys)

        if (
            value is not None
            and not isinstance(value, (dict, list, tuple))
        ):
            rows.append((label, value))

    if rows:
        _add_key_value_table(document, rows)

    _add_note(
        document,
        "This report was generated by the AI-Powered Multi-Agent "
        "Sales Intelligence System. The enabled workflow may use "
        "Python, Pandas, NumPy, Matplotlib, Streamlit and python-docx.",
    )


# ---------------------------------------------------------------------
# Word document styling
# ---------------------------------------------------------------------

def _configure_document(document: Document) -> None:
    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)

    styles = [
        ("Title", 26),
        ("Heading 1", 17),
        ("Heading 2", 12),
    ]

    for style_name, size in styles:
        try:
            style = document.styles[style_name]

            if style_name == "Heading 2":
                style.font.name = "Aptos"
            else:
                style.font.name = "Aptos Display"

            style.font.size = Pt(size)
            style.font.bold = True

        except Exception:
            pass


# ---------------------------------------------------------------------
# Public report-generation functions
# ---------------------------------------------------------------------

def generate_word_report(
    report: Any,
    output_path: str,
    dataset_name: Optional[str] = None,
) -> str:
    """Generate a Word report and return its absolute path."""
    report_dict = _normalise_report(report)

    if not report_dict:
        raise ValueError(
            "Cannot generate Word report: report data is empty."
        )

    output_path = os.path.abspath(
        os.path.normpath(output_path)
    )

    output_directory = os.path.dirname(output_path)

    if output_directory:
        os.makedirs(output_directory, exist_ok=True)

    document = Document()
    _configure_document(document)

    _add_title_page(
        document,
        report_dict,
        dataset_name,
    )

    _add_executive_summary(
        document,
        report_dict,
    )

    sections = [
        (
            "dataset_overview",
            "dataset overview",
            _add_dataset_overview,
        ),
        (
            "data_quality",
            "data quality",
            _add_data_cleaning,
        ),
        (
            "trends",
            "trend",
            _add_trends,
        ),
        (
            "rankings",
            "performance",
            _add_rankings,
        ),
        (
            "kpis",
            "kpi",
            _add_kpis,
        ),
        (
            "visualizations",
            "visualization",
            _add_visualizations,
        ),
        (
            "forecast",
            "forecast",
            _add_forecast,
        ),
        (
            "insights",
            "insight",
            _add_insights,
        ),
        (
            "recommendations",
            "recommendation",
            _add_recommendations,
        ),
    ]

    for section_id, title_hint, renderer in sections:
        section_data = _find_section(
            report_dict,
            section_id,
            title_hint,
        )

        try:
            renderer(
                document,
                _content(section_data),
            )

        except Exception as exc:
            _add_heading(
                document,
                f"Section Note: {title_hint.title()}",
                level=2,
            )

            _add_note(
                document,
                "This section could not be rendered from its supplied "
                f"data. Rendering error: {type(exc).__name__}: {exc}",
            )

    final_sections = [
        (
            _add_conclusion,
            (document, report_dict),
        ),
        (
            _add_limitations,
            (document, report_dict),
        ),
        (
            _add_technical_appendix,
            (document, report_dict),
        ),
    ]

    for renderer, arguments in final_sections:
        try:
            renderer(*arguments)

        except Exception as exc:
            _add_note(
                document,
                "A final report section could not be rendered: "
                f"{type(exc).__name__}: {exc}",
            )

    document.save(output_path)

    return output_path


def build_word_report(
    report: Any,
    output_directory: str = "reports",
    filename: str = "Sales_Intelligence_Report.docx",
    dataset_name: Optional[str] = None,
) -> str:
    """Convenience wrapper around generate_word_report()."""
    os.makedirs(
        output_directory,
        exist_ok=True,
    )

    output_path = os.path.join(
        output_directory,
        filename,
    )

    return generate_word_report(
        report,
        output_path,
        dataset_name=dataset_name,
    )


if __name__ == "__main__":
    print("Word Report Generator module loaded successfully.")